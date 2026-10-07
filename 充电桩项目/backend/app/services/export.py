"""导出服务 —— PDF 3.4 工单导出 / 3.7 台账导出 / 3.9 数据导出。

实现要点（PDF 3.7）：导出 Excel 需有加载动画与充分提示。
后端提供「任务式导出」：创建任务 -> 轮询进度 -> 下载文件，
前端据此展示加载动画与提示文案。
"""

from __future__ import annotations

import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import EXPORT_DIR
from app.core.errors import NotFoundError

# 内存态导出任务登记表（生产可替换为 Redis / Celery 任务状态）
_EXPORT_TASKS: dict[str, dict] = {}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def create_task(task_type: str, title: str, operator: str | None = None) -> dict:
    task_id = str(uuid.uuid4())
    task = {
        "task_id": task_id,
        "task_type": task_type,
        "title": title,
        "operator": operator,
        "status": "running",
        "progress": 5,
        "message": "正在准备数据…",
        "file_path": None,
        "file_name": None,
        "row_count": 0,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "finished_at": None,
        "error": None,
    }
    _EXPORT_TASKS[task_id] = task
    return task


def update_task(task_id: str, **kwargs) -> dict:
    task = _EXPORT_TASKS.get(task_id)
    if task is None:
        raise NotFoundError("导出任务不存在或已过期")
    task.update(kwargs)
    return task


def get_task(task_id: str) -> dict:
    task = _EXPORT_TASKS.get(task_id)
    if task is None:
        raise NotFoundError("导出任务不存在或已过期")
    return task


def list_tasks(limit: int = 20) -> list[dict]:
    items = sorted(_EXPORT_TASKS.values(), key=lambda t: t["created_at"], reverse=True)
    return items[:limit]


def write_excel(
    rows: list[dict],
    *,
    sheet_title: str,
    file_stem: str,
    meta_lines: list[str] | None = None,
) -> Path:
    """把字典列表写成带样式的 Excel 文件。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Sheet1"

    row_cursor = 1
    if meta_lines:
        for line in meta_lines:
            ws.cell(row=row_cursor, column=1, value=line).font = Font(bold=True, size=11)
            row_cursor += 1
        row_cursor += 1

    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if not headers:
        headers = ["提示"]
        rows = [{"提示": "当前查询条件下没有数据"}]

    header_row = row_cursor
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for r, row in enumerate(rows, start=header_row + 1):
        for col, name in enumerate(headers, start=1):
            value = row.get(name)
            if isinstance(value, (list, dict)):
                value = str(value)
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    # 列宽自适应（中文按 2 个字符宽度估算）
    for col, name in enumerate(headers, start=1):
        max_len = _display_len(str(name))
        for row in rows[:500]:
            value = row.get(name)
            if value is None:
                continue
            max_len = max(max_len, _display_len(str(value)))
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(10, max_len + 4))

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    path = EXPORT_DIR / f"{file_stem}_{timestamp}.xlsx"
    wb.save(path)
    return path


def _display_len(text: str) -> int:
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


async def export_work_orders(
    db: AsyncSession, *, task_id: str, operator: str | None = None, order_ids: list[str] | None = None
) -> dict:
    """工单导出（PDF 3.4）。"""
    from app.services.work_order import export_rows

    update_task(task_id, progress=20, message="正在读取工单数据…")
    rows = await export_rows(db, order_ids)
    update_task(task_id, progress=60, message=f"正在生成 Excel（共 {len(rows)} 行）…")
    path = write_excel(
        rows,
        sheet_title="工单明细",
        file_stem="work_orders",
        meta_lines=[
            f"工单导出报表   导出人：{operator or '-'}",
            f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   记录数：{len(rows)}",
        ],
    )
    return update_task(
        task_id,
        status="success",
        progress=100,
        message=f"导出完成，共 {len(rows)} 行",
        file_path=str(path),
        file_name=path.name,
        row_count=len(rows),
        finished_at=datetime.now().isoformat(timespec="seconds"),
    )


async def export_ledger(
    db: AsyncSession, *, task_id: str, operator: str | None = None, **filters
) -> dict:
    """台账导出（PDF 3.7：站台台账 / 充电桩台账，信息导出）。"""
    from app.services.asset import ledger_export_rows

    update_task(task_id, progress=20, message="正在读取台账数据…")
    filters.pop("task_id", None)
    filters.pop("operator", None)
    rows = await ledger_export_rows(db, **filters)
    update_task(task_id, progress=60, message=f"正在生成 Excel（共 {len(rows)} 行）…")
    path = write_excel(
        rows,
        sheet_title="资产台账",
        file_stem="asset_ledger",
        meta_lines=[
            f"资产台账导出报表   导出人：{operator or '-'}",
            f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   记录数：{len(rows)}",
        ],
    )
    return update_task(
        task_id,
        status="success",
        progress=100,
        message=f"导出完成，共 {len(rows)} 行",
        file_path=str(path),
        file_name=path.name,
        row_count=len(rows),
        finished_at=datetime.now().isoformat(timespec="seconds"),
    )


async def export_statistics(
    db: AsyncSession, *, task_id: str, operator: str | None = None, **kwargs
) -> dict:
    """统计数据导出（PDF 3.9 数据导出功能）。"""
    from app.services.statistics import dashboard_overview

    update_task(task_id, progress=20, message="正在汇总统计数据…")
    data = await dashboard_overview(db, **kwargs)
    rows: list[dict] = []

    wo = data["work_order"]
    rows.append({"统计维度": "工单总数", "数值": wo["total"]})
    rows.append({"统计维度": "待办工单", "数值": wo["pending"]})
    rows.append({"统计维度": "已办工单", "数值": wo["done"]})
    rows.append({"统计维度": "逾期工单", "数值": wo["overdue"]})
    rows.append({"统计维度": "紧急工单", "数值": wo["urgent"]})
    rows.append({"统计维度": "完成率(%)", "数值": wo["completion_rate"]})
    rows.append({"统计维度": "逾期率(%)", "数值": wo["overdue_rate"]})
    rows.append({"统计维度": "故障总数", "数值": data["fault"]["total"]})
    rows.append({"统计维度": "故障核查率(%)", "数值": data["fault"]["verify_rate"]})
    rows.append({"统计维度": "巡检记录数", "数值": data["inspection"]["total"]})
    rows.append({"统计维度": "巡检异常记录数", "数值": data["inspection"]["abnormal_records"]})

    for item in data["order_type_dist"]:
        rows.append({"统计维度": f"工单类型-{item['name']}", "数值": item["value"]})
    for item in data["project_rank"]:
        rows.append({"统计维度": f"项目排名-{item['name']}", "数值": item["value"]})
    for item in data["defect_station_rank"]:
        rows.append({"统计维度": f"站点消缺排名-{item['name']}", "数值": item["value"]})

    update_task(task_id, progress=70, message="正在生成 Excel…")
    path = write_excel(
        rows,
        sheet_title="统计分析",
        file_stem="statistics",
        meta_lines=[
            f"统计分析导出报表   导出人：{operator or '-'}",
            f"统计区间：{data['period']['start']} ~ {data['period']['end']}",
            f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        ],
    )
    return update_task(
        task_id,
        status="success",
        progress=100,
        message=f"导出完成，共 {len(rows)} 行",
        file_path=str(path),
        file_name=path.name,
        row_count=len(rows),
        finished_at=datetime.now().isoformat(timespec="seconds"),
    )
